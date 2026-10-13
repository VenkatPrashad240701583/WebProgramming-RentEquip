#!/usr/bin/env python3
"""
RentEquip - College Project Review Presentation Generator (10 Slides)
Subject: Web Programming
Student: VENKAT PRASHAD A (Roll Number: 240701583)
Department: Department of Computer Science & Engineering
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

# Dimensions: 16:9 widescreen (13.333 x 7.5 inches)
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Color Palette: Campus Academic Theme
COLOR_BG = RGBColor(252, 250, 244)          # Warm notebook cream
COLOR_PRIMARY = RGBColor(27, 67, 50)        # Deep Academic Forest Green (#1B4332)
COLOR_ORANGE = RGBColor(255, 107, 53)       # Vibrant student orange (#FF6B35)
COLOR_AMBER = RGBColor(245, 158, 11)        # Warm sunshine yellow (#F59E0B)
COLOR_BLUE = RGBColor(2, 132, 199)          # Sky blue
COLOR_TEXT_MAIN = RGBColor(28, 40, 38)      # Dark slate ink
COLOR_TEXT_MUTED = RGBColor(75, 93, 86)     # Academic green-slate
COLOR_BORDER = RGBColor(232, 226, 210)      # Lined paper border

os.makedirs('assets/presentation', exist_ok=True)

def generate_diagrams():
    """Generates crisp, high-resolution diagram graphics for slides using Pillow without emojis"""
    
    # 1. Slide 1 Graphic: Campus gear visual
    im1 = Image.new('RGB', (800, 600), '#f7f4ea')
    d1 = ImageDraw.Draw(im1)
    d1.rectangle([(20, 20), (780, 580)], fill='#ffffff', outline='#e8e2d2', width=3)
    
    # Box 1: Stationery
    d1.rounded_rectangle([(50, 50), (370, 270)], radius=12, fill='#fef3c7', outline='#f59e0b', width=2)
    d1.text((70, 75), "Stationery & Tools", fill='#92400e', font_size=24)
    d1.text((70, 115), "- Pilot V7 Pens & Pencils\n- Geometry Boxes & Sets\n- Compasses & Rulers\n- Whiteboard Markers", fill='#78350f', font_size=18)
    
    # Box 2: Drawing Tools
    d1.rounded_rectangle([(410, 50), (730, 270)], radius=12, fill='#e0f2fe', outline='#0284c7', width=2)
    d1.text((430, 75), "Engineering Drawing", fill='#0369a1', font_size=24)
    d1.text((430, 115), "- Pine Wood Drawing Boards\n- Technical Mini-Drafters\n- French Curve Sets\n- Technical Ink Pens", fill='#075985', font_size=18)

    # Box 3: Lab Essentials
    d1.rounded_rectangle([(50, 310), (370, 530)], radius=12, fill='#dcfce7', outline='#16a34a', width=2)
    d1.text((70, 335), "Lab Essentials", fill='#166534', font_size=24)
    d1.text((70, 375), "- Pure Cotton Lab Coats\n- Chemical Safety Goggles\n- Digital Multimeters\n- Arduino Prototyping Kits", fill='#14532d', font_size=18)

    # Box 4: Academic Resources
    d1.rounded_rectangle([(410, 310), (730, 530)], radius=12, fill='#ede9fe', outline='#7c3aed', width=2)
    d1.text((430, 335), "Calculators & Books", fill='#5b21b6', font_size=24)
    d1.text((430, 375), "- Casio fx-991EX Calculators\n- TI-84 Plus Graphing Tools\n- Semester Math Textbooks\n- GATE CSE Solved Papers", fill='#4c1d95', font_size=18)
    im1.save('assets/presentation/slide1_campus.png')

    # 2. Slide 2 Graphic: Problem Statement Infographic
    im2 = Image.new('RGB', (800, 520), '#ffffff')
    d2 = ImageDraw.Draw(im2)
    problems = [
        ("Forgotten Academic Items", "Students arrive at labs or exams missing required pens, calculators, or lab coats.", "#fee2e2", "#b91c1c"),
        ("High Cost of Single-Use Purchases", "Paying full price for specialized tools used only for one semester is wasteful.", "#fef3c7", "#b45309"),
        ("Lack of Immediate Campus Access", "No organized place on campus to quickly borrow gear during emergencies.", "#ede9fe", "#6d28d9"),
        ("Unused Items Sitting in Hostels", "Senior students possess working equipment with no platform to lend to juniors.", "#e0f2fe", "#0369a1"),
    ]
    for i, (title, desc, bg, fg) in enumerate(problems):
        y = 30 + i * 115
        d2.rounded_rectangle([(30, y), (770, y + 95)], radius=10, fill=bg, outline=fg, width=2)
        d2.text((50, y + 16), title, fill=fg, font_size=22)
        d2.text((50, y + 52), desc, fill='#334155', font_size=16)
    im2.save('assets/presentation/slide2_problem.png')

    # 3. Slide 3 Graphic: Solution Concept
    im3 = Image.new('RGB', (800, 520), '#ffffff')
    d3 = ImageDraw.Draw(im3)
    sol_cards = [
        ("Student-to-Student Sharing", "Direct connection between student borrowers and student equipment owners.", "#dcfce7", "#15803d"),
        ("Centralized Campus Hub", "One web portal covering stationery, drawing boards, electronics, and books.", "#e0f2fe", "#0284c7"),
        ("Low Daily Rental Rates", "Affordable rates in Indian Rupees (starting from Rs 2/day) with refundable deposits.", "#fef3c7", "#b45309"),
        ("Trusted Peer Community", "Transparent item conditions, campus pickup spots, and verified peer feedback.", "#ede9fe", "#7c3aed"),
    ]
    for i, (title, desc, bg, fg) in enumerate(sol_cards):
        y = 30 + i * 115
        d3.rounded_rectangle([(30, y), (770, y + 95)], radius=10, fill=bg, outline=fg, width=2)
        d3.text((50, y + 16), title, fill=fg, font_size=22)
        d3.text((50, y + 52), desc, fill='#334155', font_size=16)
    im3.save('assets/presentation/slide3_solution.png')

    # 4. Slide 4 Graphic: 6 Key Features Grid
    im4 = Image.new('RGB', (800, 500), '#fdfbf7')
    d4 = ImageDraw.Draw(im4)
    feats = [
        ("Category Search & Filters", "Quick filtering by price, subject, location, and immediate availability.", "#dcfce7", "#166534"),
        ("Live Rental Cost Calculator", "Dynamic calculation of rental days, community fee, and security deposit.", "#fef3c7", "#92400e"),
        ("Request & Approval Flow", "Submit rental requests with instant alerts to equipment owners.", "#e0f2fe", "#0369a1"),
        ("Campus Pickup Locations", "Designated handover points: Central Library, SAC Desk, Labs, and Hostels.", "#ede9fe", "#5b21b6"),
        ("Return Condition Inspection", "Fair assessment of returned items before releasing deposit refunds.", "#fce7f3", "#9d174d"),
        ("Verified Peer Reviews", "Honest ratings allowed exclusively after completed rentals.", "#ccfbf1", "#0f766e"),
    ]
    for i, (title, desc, bg, fg) in enumerate(feats):
        x = 30 + (i % 2) * 385
        y = 30 + (i // 2) * 150
        d4.rounded_rectangle([(x, y), (x + 360, y + 125)], radius=12, fill=bg, outline=fg, width=2)
        d4.text((x + 18, y + 18), title, fill=fg, font_size=20)
        d4.text((x + 18, y + 55), desc, fill='#334155', font_size=15)
    im4.save('assets/presentation/slide4_features.png')

    # 5. Slide 5 Graphic: Visual Flowchart Diagram (Requested prominently!)
    im5 = Image.new('RGB', (1100, 480), '#ffffff')
    d5 = ImageDraw.Draw(im5)
    d5.rectangle([(10, 10), (1090, 470)], fill='#faf8f2', outline='#e8e2d2', width=2)
    
    # Main path boxes
    steps = [
        ("1. Search Item", "Browse catalog or filter\nby department & price", 40, 60),
        ("2. Select Dates", "Choose pickup & return\nCheck availability", 250, 60),
        ("3. Send Request", "Submit borrow request\nDeposit placed in escrow", 460, 60),
        ("4. Owner Decision", "Host reviews request\nAccepts or declines", 670, 60),
    ]
    for title, sub, x, y in steps:
        d5.rounded_rectangle([(x, y), (x + 175, y + 95)], radius=8, fill='#1b4332', outline='#0f2b20', width=2)
        d5.text((x + 12, y + 14), title, fill='#ffffff', font_size=18)
        d5.text((x + 12, y + 42), sub, fill='#d1dcd7', font_size=12)
        # Arrow right
        if x < 670:
            d5.polygon([(x + 185, y + 45), (x + 205, y + 45), (x + 205, y + 40), (x + 225, y + 47), (x + 205, y + 55), (x + 205, y + 50), (x + 185, y + 50)], fill='#ff6b35')

    # Branch 1: REJECTED (Red branch)
    d5.line([(757, 155), (757, 210)], fill='#dc2626', width=3)
    d5.polygon([(750, 210), (765, 210), (757, 225)], fill='#dc2626')
    d5.rounded_rectangle([(670, 225), (845, 305)], radius=8, fill='#fee2e2', outline='#dc2626', width=2)
    d5.text((685, 235), "Request Declined", fill='#b91c1c', font_size=17)
    d5.text((685, 260), "Full deposit refunded\nProcess ends safely", fill='#7f1d1d', font_size=13)

    # Branch 2: APPROVED (Green Path continuing downward and left)
    d5.line([(845, 107), (900, 107), (900, 225)], fill='#16a34a', width=3)
    d5.polygon([(895, 225), (905, 225), (900, 235)], fill='#16a34a')
    d5.rounded_rectangle([(875, 235), (1050, 315)], radius=8, fill='#dcfce7', outline='#16a34a', width=2)
    d5.text((890, 245), "5. Handover Pickup", fill='#15803d', font_size=17)
    d5.text((890, 270), "Student collects item\nStatus becomes Active", fill='#14532d', font_size=13)

    # Path to return, inspect, and review
    flow_bottom = [
        ("8. Leave Review", "Verified peer rating\nbuilds community trust", 40, 350, '#ede9fe', '#6d28d9'),
        ("7. Settle Deposit", "Fair deductions if any\nEscrow refund released", 320, 350, '#fef3c7', '#b45309'),
        ("6. Return & Check", "Item returned to host\nCondition is verified", 600, 350, '#e0f2fe', '#0369a1'),
    ]
    # Arrow down from pickup to return
    d5.line([(962, 315), (962, 385), (785, 385)], fill='#0284c7', width=3)
    d5.polygon([(785, 380), (785, 390), (775, 385)], fill='#0284c7')

    for title, sub, x, y, bg, fg in flow_bottom:
        d5.rounded_rectangle([(x, y), (x + 185, y + 80)], radius=8, fill=bg, outline=fg, width=2)
        d5.text((x + 12, y + 12), title, fill=fg, font_size=17)
        d5.text((x + 12, y + 36), sub, fill='#334155', font_size=12)

    # Arrow from 6 to 7
    d5.polygon([(590, 387), (570, 387), (570, 382), (550, 390), (570, 398), (570, 393), (590, 393)], fill='#0284c7')
    # Arrow from 7 to 8
    d5.polygon([(310, 387), (290, 387), (290, 382), (270, 390), (290, 398), (290, 393), (310, 393)], fill='#0284c7')
    im5.save('assets/presentation/slide5_flowchart.png')

    # 6. Slide 6 Graphic: Three User Roles
    im6 = Image.new('RGB', (800, 500), '#fdfbf7')
    d6 = ImageDraw.Draw(im6)
    roles = [
        ("Student Renter (Borrower)", "Aarav Sharma - 2nd Year CSE\n- Searches supplies across categories\n- Requests short-term rentals\n- Confirms pickup & tracks returns\n- Submits verified peer feedback", "#dcfce7", "#15803d"),
        ("Equipment Owner (Lender)", "Priya Patel - Final Year Mech\n- Lists spare tools & instruments\n- Approves or declines borrow requests\n- Inspects items upon return\n- Releases security deposit refunds", "#e0f2fe", "#0369a1"),
        ("Campus Administrator", "SAC Desk / Central Library\n- Oversees platform activity & users\n- Audits equipment and borrow logs\n- Enforces emergency status overrides\n- Manages data backups and resets", "#ede9fe", "#6d28d9"),
    ]
    for i, (title, desc, bg, fg) in enumerate(roles):
        y = 30 + i * 150
        d6.rounded_rectangle([(30, y), (770, y + 125)], radius=12, fill=bg, outline=fg, width=2)
        d6.text((50, y + 16), title, fill=fg, font_size=22)
        d6.text((50, y + 50), desc, fill='#334155', font_size=15)
    im6.save('assets/presentation/slide6_roles.png')

    # 7. Slide 7 Graphic: Simple Tech Stack Architecture
    im7 = Image.new('RGB', (800, 500), '#ffffff')
    d7 = ImageDraw.Draw(im7)
    stack_boxes = [
        ("HTML5 (Structure)", "Builds all semantic pages, input forms, and accessible modal dialogs.", "#fee2e2", "#b91c1c"),
        ("CSS3 (Styling & Animation)", "Custom campus notebook theme, responsive grid layouts, and smooth animations.", "#e0f2fe", "#0369a1"),
        ("Vanilla JavaScript (ES6+)", "Client-side rental calculations, date conflict engine, and modal interactions.", "#fef3c7", "#b45309"),
        ("Browser localStorage", "Saves all demonstration data, bookings, users, and reviews directly in browser.", "#dcfce7", "#15803d"),
        ("GitHub & GitHub Actions", "Source code repository and automated CI/CD pipeline on every push to main.", "#ede9fe", "#6d28d9"),
        ("GitHub Pages (Hosting)", "Deploys static files directly to the web with zero server maintenance.", "#ccfbf1", "#0f766e"),
    ]
    for i, (title, desc, bg, fg) in enumerate(stack_boxes):
        x = 30 + (i % 2) * 385
        y = 30 + (i // 2) * 150
        d7.rounded_rectangle([(x, y), (x + 360, y + 125)], radius=12, fill=bg, outline=fg, width=2)
        d7.text((x + 18, y + 18), title, fill=fg, font_size=20)
        d7.text((x + 18, y + 52), desc, fill='#334155', font_size=15)
    im7.save('assets/presentation/slide7_tech.png')

    # 8. Slide 8 Graphic: Deployment & Access
    im8 = Image.new('RGB', (800, 500), '#f8fafc')
    d8 = ImageDraw.Draw(im8)
    deploy_points = [
        ("Zero Installation", "Runs directly inside any modern web browser without installing software.", "#dcfce7", "#166534"),
        ("Public Web Access", "Hosted online through GitHub Pages URL accessible to anyone.", "#e0f2fe", "#0284c7"),
        ("Fully Responsive", "Adapts seamlessly across mobile phones, tablets, and desktop displays.", "#fef3c7", "#92400e"),
        ("Continuous Deployment", "Pushing updates to GitHub main branch automatically redeploys site.", "#ede9fe", "#5b21b6"),
    ]
    for i, (title, desc, bg, fg) in enumerate(deploy_points):
        y = 30 + i * 115
        d8.rounded_rectangle([(30, y), (770, y + 95)], radius=10, fill=bg, outline=fg, width=2)
        d8.text((50, y + 16), title, fill=fg, font_size=22)
        d8.text((50, y + 52), desc, fill='#334155', font_size=16)
    im8.save('assets/presentation/slide8_deploy.png')

    # 9. Slide 9 Graphic: Benefits & Future Scope
    im9 = Image.new('RGB', (800, 500), '#ffffff')
    d9 = ImageDraw.Draw(im9)
    # Left: Benefits
    d9.rounded_rectangle([(30, 30), (380, 470)], radius=14, fill='#dcfce7', outline='#16a34a', width=2)
    d9.text((50, 50), "Immediate Benefits", fill='#166534', font_size=22)
    d9.text((50, 95), "- Easy access to academic items\n- Major savings for students\n- Extra income from idle gear\n- Reduces plastic & tool waste\n- Stronger campus community", fill='#14532d', font_size=16)

    # Right: Future Scope
    d9.rounded_rectangle([(420, 30), (770, 470)], radius=14, fill='#e0f2fe', outline='#0284c7', width=2)
    d9.text((440, 50), "Future Scope", fill='#0369a1', font_size=22)
    d9.text((440, 95), "- Real campus UPI integration\n- College ID card verification\n- QR code smart locker pickups\n- Multi-college network sharing\n- Demand & supply analytics", fill='#0c4a6e', font_size=16)
    im9.save('assets/presentation/slide9_scope.png')

    # 10. Slide 10 Graphic: Conclusion Card
    im10 = Image.new('RGB', (800, 500), '#1b4332')
    d10 = ImageDraw.Draw(im10)
    d10.rectangle([(20, 20), (780, 480)], outline='#ff6b35', width=3)
    d10.text((60, 60), "RentEquip Summary", fill='#f59e0b', font_size=28)
    d10.text((60, 120), "• Connects student lenders with borrowers\n• Solves urgent shortages of exam and lab supplies\n• Proves that clean, native web code works without bloated frameworks\n• Ready for campus demonstration and peer collaboration", fill='#ffffff', font_size=18)
    
    d10.rounded_rectangle([(60, 330), (740, 420)], radius=12, fill='#ff6b35')
    d10.text((100, 355), "\"Rent smarter. Share better. Learn together.\"", fill='#ffffff', font_size=24)
    im10.save('assets/presentation/slide10_conclusion.png')

def create_presentation():
    generate_diagrams()
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank_layout = prs.slide_layouts[6]

    def add_standard_header(slide, title_text, category_text, slide_number_text):
        # Header background
        header_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.15)
        )
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = COLOR_PRIMARY
        header_shape.line.fill.background()

        # Orange accent stripe
        accent_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), SLIDE_WIDTH, Inches(0.06)
        )
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = COLOR_ORANGE
        accent_line.line.fill.background()

        # Category text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(6), Inches(0.28))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_AMBER

        # Title text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.5), Inches(0.65))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255)

        # Slide number
        num_box = slide.shapes.add_textbox(Inches(11.2), Inches(0.35), Inches(1.5), Inches(0.5))
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = slide_number_text
        p_num.font.size = Pt(16)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_AMBER
        p_num.alignment = PP_ALIGN.RIGHT

    # -------------------------------------------------------------
    # SLIDE 1: Title
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    hero_left = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(6.8), SLIDE_HEIGHT
    )
    hero_left.fill.solid()
    hero_left.fill.fore_color.rgb = COLOR_PRIMARY
    hero_left.line.fill.background()

    hero_right = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(0), Inches(6.533), SLIDE_HEIGHT
    )
    hero_right.fill.solid()
    hero_right.fill.fore_color.rgb = COLOR_BG
    hero_right.line.fill.background()

    tag_box = slide1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(5.5), Inches(0.4))
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "COLLEGE PROJECT PRESENTATION"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_AMBER

    pname_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(5.5), Inches(1.2))
    tf_pname = pname_box.text_frame
    p_pname = tf_pname.paragraphs[0]
    p_pname.text = "RentEquip"
    p_pname.font.size = Pt(50)
    p_pname.font.bold = True
    p_pname.font.color.rgb = RGBColor(255, 255, 255)

    sub_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(5.5), Inches(0.9))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Campus Equipment Rental Marketplace"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(209, 220, 215)

    div = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.45), Inches(4.5), Inches(0.05)
    )
    div.fill.solid()
    div.fill.fore_color.rgb = COLOR_ORANGE
    div.line.fill.background()

    cred_box = slide1.shapes.add_textbox(Inches(0.8), Inches(3.7), Inches(5.5), Inches(3.2))
    tf_cred = cred_box.text_frame
    tf_cred.word_wrap = True

    p_c1 = tf_cred.paragraphs[0]
    p_c1.text = "Course: Web Programming"
    p_c1.font.size = Pt(16)
    p_c1.font.bold = True
    p_c1.font.color.rgb = RGBColor(254, 243, 199)

    p_c2 = tf_cred.add_paragraph()
    p_c2.space_before = Pt(8)
    p_c2.text = "Student Name: VENKAT PRASHAD A"
    p_c2.font.size = Pt(20)
    p_c2.font.bold = True
    p_c2.font.color.rgb = RGBColor(255, 255, 255)

    p_c3 = tf_cred.add_paragraph()
    p_c3.space_before = Pt(4)
    p_c3.text = "Roll Number: 240701583"
    p_c3.font.size = Pt(18)
    p_c3.font.bold = True
    p_c3.font.color.rgb = COLOR_AMBER

    p_c4 = tf_cred.add_paragraph()
    p_c4.space_before = Pt(4)
    p_c4.text = "Department of Computer Science & Engineering"
    p_c4.font.size = Pt(14)
    p_c4.font.color.rgb = RGBColor(163, 184, 176)

    slide1.shapes.add_picture('assets/presentation/slide1_campus.png', Inches(7.2), Inches(0.9), Inches(5.5), Inches(5.5))

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_standard_header(slide2, "Problem Statement", "The Campus Challenge", "02 / 10")

    box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))
    tf2 = box2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "Challenges College Students Face:"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets2 = [
        ("Forgotten Essential Items", "Students frequently forget items like pens, calculators, lab coats, and drawing tools right before tests or practicals."),
        ("High One-Time Purchase Costs", "Buying equipment like drawing boards or scientific calculators for a single semester is expensive and wasteful."),
        ("Temporary Academic Needs", "Students struggle to find specialized tools needed only for short-term project presentations or lab assignments."),
        ("Lack of Centralized Platform", "There is no organized, convenient campus marketplace for students to easily borrow or rent items from peers.")
    ]
    for b_title, b_desc in bullets2:
        p_b = tf2.add_paragraph()
        p_b.space_before = Pt(12)
        r1 = p_b.add_run()
        r1.text = "• " + b_title + ": "
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p_b.add_run()
        r2.text = b_desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_TEXT_MUTED

    slide2.shapes.add_picture('assets/presentation/slide2_problem.png', Inches(7.1), Inches(1.5), Inches(5.5), Inches(5.3))

    # -------------------------------------------------------------
    # SLIDE 3: Proposed Solution
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_standard_header(slide3, "Proposed Solution — RentEquip", "Project Overview", "03 / 10")

    box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))
    tf3 = box3.text_frame
    tf3.word_wrap = True

    p = tf3.paragraphs[0]
    p.text = "How RentEquip Solves the Problem:"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets3 = [
        ("Peer-to-Peer Campus Marketplace", "Connects students who need academic items with students, seniors, and campus clubs who own them."),
        ("Search & Rent with Ease", "Students can browse, search, and rent essential items for flexible daily periods at affordable student rates."),
        ("Lender Tools for Owners", "Equipment owners can list spare items, set rental prices and deposits, and manage incoming requests easily."),
        ("Single Centralized Platform", "A unified website simplifies discovery, date reservation, pickup coordination, and return settlements.")
    ]
    for b_title, b_desc in bullets3:
        p_b = tf3.add_paragraph()
        p_b.space_before = Pt(12)
        r1 = p_b.add_run()
        r1.text = "• " + b_title + ": "
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p_b.add_run()
        r2.text = b_desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_TEXT_MUTED

    slide3.shapes.add_picture('assets/presentation/slide3_solution.png', Inches(7.1), Inches(1.5), Inches(5.5), Inches(5.3))

    # -------------------------------------------------------------
    # SLIDE 4: Key Features
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_standard_header(slide4, "Key Features", "What Students Can Do", "04 / 10")

    box4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))
    tf4 = box4.text_frame
    tf4.word_wrap = True

    p = tf4.paragraphs[0]
    p.text = "Core Application Capabilities:"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets4 = [
        ("Smart Search & Filters", "Filter equipment by category, daily price budget, availability, and campus pickup locations."),
        ("Diverse Academic Inventory", "Browse stationery, engineering drawing boards, lab coats, multimeters, tech gear, art sets, and textbooks."),
        ("Seamless Rental Requests", "Select pickup and return dates, view live cost quotes, and send booking requests with a single click."),
        ("Owner Request Management", "Hosts can review incoming requests and choose to approve or decline with feedback."),
        ("Deposit & Return Tracking", "Transparent tracking of refundable security deposits, actual return dates, and condition reports."),
        ("Alerts & Verified Reviews", "Real-time notifications for booking updates and verified peer reviews for completed rentals.")
    ]
    for b_title, b_desc in bullets4:
        p_b = tf4.add_paragraph()
        p_b.space_before = Pt(8)
        r1 = p_b.add_run()
        r1.text = "• " + b_title + ": "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p_b.add_run()
        r2.text = b_desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = COLOR_TEXT_MUTED

    slide4.shapes.add_picture('assets/presentation/slide4_features.png', Inches(7.1), Inches(1.5), Inches(5.5), Inches(5.3))

    # -------------------------------------------------------------
    # SLIDE 5: System Workflow (Visual Flowchart!)
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_standard_header(slide5, "System Workflow", "End-to-End Rental Journey", "05 / 10")

    # Workflow narrative header
    wf_box = slide5.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.8), Inches(0.7))
    tf_wf = wf_box.text_frame
    tf_wf.word_wrap = True
    p_wf = tf_wf.paragraphs[0]
    p_wf.text = "How RentEquip Works: From Search to Return & Review"
    p_wf.font.size = Pt(18)
    p_wf.font.bold = True
    p_wf.font.color.rgb = COLOR_PRIMARY

    # Full visual flowchart diagram image on Slide 5
    slide5.shapes.add_picture('assets/presentation/slide5_flowchart.png', Inches(0.8), Inches(2.1), Inches(11.7), Inches(5.0))

    # -------------------------------------------------------------
    # SLIDE 6: User Roles
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_standard_header(slide6, "User Roles & Ecosystem", "Platform Stakeholders", "06 / 10")

    box6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))
    tf6 = box6.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "Three Distinct Campus Roles:"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets6 = [
        ("1. Customer (Student Borrower)", "Searches campus catalog, selects rental duration, submits borrow requests, tracks bookings, confirms pickup, and writes peer reviews."),
        ("2. Equipment Owner (Student Lender / Senior)", "Lists spare academic items, sets daily prices and deposits, approves or declines requests, checks returned items, and releases deposits."),
        ("3. Administrator (Student Affairs Council / SAC)", "Oversees platform activity, user registry, campus catalog audit, monitors overdue alerts, and maintains data backups.")
    ]
    for b_title, b_desc in bullets6:
        p_b = tf6.add_paragraph()
        p_b.space_before = Pt(14)
        r1 = p_b.add_run()
        r1.text = b_title + "\n"
        r1.font.bold = True
        r1.font.size = Pt(15)
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p_b.add_run()
        r2.text = b_desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_TEXT_MUTED

    slide6.shapes.add_picture('assets/presentation/slide6_roles.png', Inches(7.1), Inches(1.5), Inches(5.5), Inches(5.3))

    # -------------------------------------------------------------
    # SLIDE 7: Technology Used
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_standard_header(slide7, "Technology Used", "Simple Web Architecture", "07 / 10")

    box7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))
    tf7 = box7.text_frame
    tf7.word_wrap = True

    p = tf7.paragraphs[0]
    p.text = "Built with Core Web Technologies:"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets7 = [
        ("HTML", "Builds the web pages, accessible forms, modal structures, and semantic layouts."),
        ("CSS", "Provides the campus notebook visual design, responsive layouts (Flexbox & Grid), and smooth animations."),
        ("JavaScript (Vanilla ES6)", "Handles user interactions, real-time rental calculations, date conflict checking, and page navigation."),
        ("Browser localStorage", "Saves all demonstration data (equipment, bookings, users, and reviews) directly in the browser with zero backend."),
        ("GitHub & GitHub Actions", "Hosts source code repository and runs automated CI/CD deployment on every push to main."),
        ("GitHub Pages", "Deploys and hosts the static website online with a publicly accessible URL.")
    ]
    for b_title, b_desc in bullets7:
        p_b = tf7.add_paragraph()
        p_b.space_before = Pt(8)
        r1 = p_b.add_run()
        r1.text = "• " + b_title + ": "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p_b.add_run()
        r2.text = b_desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = COLOR_TEXT_MUTED

    slide7.shapes.add_picture('assets/presentation/slide7_tech.png', Inches(7.1), Inches(1.5), Inches(5.5), Inches(5.3))

    # -------------------------------------------------------------
    # SLIDE 8: Deployment & Accessibility
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_standard_header(slide8, "Deployment & Accessibility", "Hosting & Delivery", "08 / 10")

    box8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))
    tf8 = box8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "How Users Access RentEquip:"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets8 = [
        ("Hosted on GitHub Pages", "The website is deployed live using GitHub Pages infrastructure, ensuring high availability."),
        ("Accessible via Public URL", "Students and faculty can access the marketplace directly through a standard web link."),
        ("No Installation Required", "Zero client setup needed — runs immediately in any modern desktop or mobile web browser."),
        ("Cross-Device Responsive", "Clean adaptive layouts tailored for smartphones, tablets, and laptops."),
        ("Automated Updates", "Pushing updates to the main branch automatically triggers GitHub Actions to publish changes live.")
    ]
    for b_title, b_desc in bullets8:
        p_b = tf8.add_paragraph()
        p_b.space_before = Pt(12)
        r1 = p_b.add_run()
        r1.text = "• " + b_title + ": "
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p_b.add_run()
        r2.text = b_desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_TEXT_MUTED

    slide8.shapes.add_picture('assets/presentation/slide8_deploy.png', Inches(7.1), Inches(1.5), Inches(5.5), Inches(5.3))

    # -------------------------------------------------------------
    # SLIDE 9: Benefits and Future Scope
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_standard_header(slide9, "Benefits & Future Scope", "Project Impact & Evolution", "09 / 10")

    box9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))
    tf9 = box9.text_frame
    tf9.word_wrap = True

    p = tf9.paragraphs[0]
    p.text = "Impact on Campus Community:"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets9 = [
        ("Convenient Item Access: ", "Makes essential study tools and lab gear quick and easy to locate right when needed."),
        ("Affordable Reuse: ", "Reduces unnecessary purchases and promotes sustainable sharing among students."),
        ("Earnings for Lenders: ", "Helps students and clubs earn pocket money from equipment they already own."),
        ("Future: Real UPI Payments: ", "Integration with live college payment gateways (Google Pay, PhonePe, Paytm)."),
        ("Future: Student ID Verification: ", "Campus roll number authentication and digital college ID verification."),
        ("Future: Smart Pickup Lockers: ", "QR-code enabled pickup lockers at central library and hostel entrances.")
    ]
    for b_title, b_desc in bullets9:
        p_b = tf9.add_paragraph()
        p_b.space_before = Pt(8)
        r1 = p_b.add_run()
        r1.text = "• " + b_title
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p_b.add_run()
        r2.text = b_desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = COLOR_TEXT_MUTED

    slide9.shapes.add_picture('assets/presentation/slide9_scope.png', Inches(7.1), Inches(1.5), Inches(5.5), Inches(5.3))

    # -------------------------------------------------------------
    # SLIDE 10: Conclusion
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_standard_header(slide10, "Conclusion", "Project Summary", "10 / 10")

    box10 = slide10.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))
    tf10 = box10.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.text = "Project Takeaways:"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    bullets10 = [
        ("Campus-Focused Utility", "RentEquip provides a practical, user-friendly rental marketplace tailored to actual student needs."),
        ("Connecting Students & Clubs", "Bridges the gap between students with idle equipment and peers who need temporary access."),
        ("Practical Web Engineering", "Demonstrates how core HTML, CSS, and JavaScript can deliver a responsive, interactive web application deployed live on the cloud."),
        ("Sustainable Academic Sharing", "Promotes an organized, affordable circular economy within the college campus community.")
    ]
    for b_title, b_desc in bullets10:
        p_b = tf10.add_paragraph()
        p_b.space_before = Pt(12)
        r1 = p_b.add_run()
        r1.text = "• " + b_title + ": "
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_TEXT_MAIN
        r2 = p_b.add_run()
        r2.text = b_desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom motto box
    motto_shape = slide10.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.9), Inches(6.0), Inches(0.9)
    )
    motto_shape.fill.solid()
    motto_shape.fill.fore_color.rgb = COLOR_ORANGE
    motto_shape.line.fill.background()
    tf_motto = motto_shape.text_frame
    p_m = tf_motto.paragraphs[0]
    p_m.text = "Rent smarter. Share better. Learn together."
    p_m.font.size = Pt(16)
    p_m.font.bold = True
    p_m.font.color.rgb = RGBColor(255, 255, 255)
    p_m.alignment = PP_ALIGN.CENTER

    slide10.shapes.add_picture('assets/presentation/slide10_conclusion.png', Inches(7.1), Inches(1.5), Inches(5.5), Inches(5.3))

    # Save presentation
    output_path = 'RentEquip_Presentation.pptx'
    prs.save(output_path)
    print(f"10-Slide presentation successfully generated and saved to {output_path}")

if __name__ == '__main__':
    create_presentation()
